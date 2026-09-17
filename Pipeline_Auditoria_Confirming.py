# -*- coding: utf-8 -*-
"""
Pipeline de auditoría de confirming ejecutado en GitHub Actions.

El programa procesa los documentos de confirming almacenados en Google Drive.
Identifica el banco, la sociedad y la fecha, extrae la información de los
documentos y la procesa para llevar a cabo la auditoría, genera los archivos
de salida y actualiza la base de datos. Al finalizar, crea un PDF
de monitorización con el resultado de la ejecución.

La API de Google Drive permite subir, consultar y descargar los archivos del
proyecto. Para minimizar la intervención manual, GitHub Actions ejecuta
periódicamente el proceso generando la información necesaria para la auditoría
de los documentos y guardando la información en una base de datos en.
"""
import os
import re
import csv
import subprocess
from datetime import datetime
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

from pydrive2.auth import GoogleAuth
from pydrive2.drive import GoogleDrive
from oauth2client.client import OAuth2Credentials

# CONEXIÓN CON GOOGLE DRIVE
"""
Esta parte utiliza las credenciales OAuth guardadas como Secrets del repositorio.
El refresh token, obtenido previamente con obtener_refresh_token.py, permite
renovar el acceso a Drive sin iniciar sesión en cada ejecución.
"""

# Nombre de la carpeta de Drive que contiene todos los elementos del proyecto.
NOMBRE_CARPETA_RAIZ = "Bonificaciones Confirming"


# Crea la conexión con Google Drive a partir de los Secrets de GitHub.
def conectar_drive():
    """
    La función recoge el identificador del cliente, su clave y el refresh
    token desde variables de entorno. Después entrega estas credenciales a
    PyDrive2, que puede solicitar automáticamente nuevos tokens de acceso.
    """
    credenciales = OAuth2Credentials(
        access_token=None,
        client_id=os.environ["GDRIVE_CLIENT_ID"],
        client_secret=os.environ["GDRIVE_CLIENT_SECRET"],
        refresh_token=os.environ["GDRIVE_REFRESH_TOKEN"],
        token_expiry=None,
        token_uri="https://oauth2.googleapis.com/token",
        user_agent="pipeline-confirming-actions/1.0",
    )
    gauth = GoogleAuth()
    gauth.credentials = credenciales
    return GoogleDrive(gauth)


drive = conectar_drive()

# GESTIÓN DE CARPETAS Y ARCHIVOS EN DRIVE
"""
Drive identifica cada elemento mediante un id y relaciona las carpetas a
través del campo "parents". Las siguientes funciones traducen los nombres
utilizados en el proyecto a esos identificadores internos y centralizan
las operaciones de búsqueda, creación, subida y movimiento de archivos.
"""

# Busca en Drive una carpeta que coincida con el nombre recibido.
def buscar_carpeta_raiz(nombre):
    """
    La consulta limita los resultados a carpetas que no estén en la
    papelera. Si existe alguna coincidencia devuelve la primera; en caso
    contrario devuelve None para que el programa pueda detenerse.
    """
    consulta = f"title='{nombre}' and mimeType='application/vnd.google-apps.folder' and trashed=false"
    resultados = drive.ListFile({"q": consulta}).GetList()
    return resultados[0] if resultados else None


# Localiza una subcarpeta y la crea cuando todavía no existe.
def buscar_o_crear_carpeta(nombre, id_padre):
    """
    La búsqueda se realiza dentro del id de la carpeta padre. La función
    devuelve siempre un identificador válido, tanto si encuentra la carpeta
    como si necesita crearla durante esa misma ejecución.
    """
    consulta = (f"'{id_padre}' in parents and title='{nombre}' "
                f"and mimeType='application/vnd.google-apps.folder' and trashed=false")
    resultados = drive.ListFile({"q": consulta}).GetList()
    if resultados:
        return resultados[0]["id"]
    carpeta_nueva = drive.CreateFile({
        "title": nombre,
        "mimeType": "application/vnd.google-apps.folder",
        "parents": [{"id": id_padre}],
    })
    carpeta_nueva.Upload()
    return carpeta_nueva["id"]


# Obtiene la carpeta correspondiente a una combinación sociedad-banco.
def carpeta_banco(banco, sociedad):
    """
    Primero localiza o crea la carpeta de la sociedad y, dentro de ella,
    localiza o crea la del banco. Así se reutiliza la misma jerarquía al
    guardar los resultados y al archivar los documentos originales.
    """
    return buscar_o_crear_carpeta(banco, buscar_o_crear_carpeta(sociedad, ID_RAIZ))


# Devuelve los archivos activos que contiene una carpeta de Drive.
def listar_archivos(id_carpeta):
    """
    La consulta excluye los elementos enviados a la papelera. Se utiliza
    para localizar documentos pendientes y comprobar si ya existen el
    histórico o el informe diario que debe actualizarse.
    """
    consulta = f"'{id_carpeta}' in parents and trashed=false"
    return drive.ListFile({"q": consulta}).GetList()


# Sube a Drive un archivo generado temporalmente por el programa.
def subir_archivo(ruta_local, nombre, id_carpeta):
    """
    La función crea el archivo en la carpeta indicada, asigna el nombre de
    salida y copia el contenido desde su ruta local. Los archivos temporales
    se encuentran en /tmp durante la ejecución de GitHub Actions.
    """
    archivo = drive.CreateFile({"title": nombre, "parents": [{"id": id_carpeta}]})
    archivo.SetContentFile(ruta_local)
    archivo.Upload()


# Sustituye el informe anterior del mismo día por la versión más reciente.
def subir_o_actualizar_archivo(ruta_local, nombre, id_carpeta):
    """
    Primero se sube el PDF nuevo y después se envían a la papelera las copias
    anteriores que tengan el mismo nombre. Este orden evita perder el informe
    si la nueva subida falla y reduce problemas con la vista previa de Drive.
    """
    anteriores = [a for a in listar_archivos(id_carpeta) if a["title"] == nombre]
    subir_archivo(ruta_local, nombre, id_carpeta)
    for anterior in anteriores:
        anterior.Trash()


# Traslada un archivo de Drive a otra carpeta.
def mover_archivo(archivo_drive, id_carpeta_destino):
    """
    El movimiento se realiza sustituyendo el identificador almacenado en
    el campo "parents" y actualizando el archivo. El resultado equivale a
    moverlo entre carpetas en un sistema de archivos convencional.
    """
    archivo_drive["parents"] = [{"id": id_carpeta_destino}]
    archivo_drive.Upload()

# PREPARACIÓN DE LA ESTRUCTURA DE DRIVE
"""
Antes de procesar documentos se comprueba que la carpeta raíz existe.
También se obtienen los identificadores de la carpeta de entrada y del
historial. Si la raíz no aparece, el programa se detiene con un mensaje
que permite detectar un nombre incorrecto o una falta de acceso.
"""
id_raiz = buscar_carpeta_raiz(NOMBRE_CARPETA_RAIZ)
if id_raiz is None:
    raise RuntimeError(
        f"No se encuentra la carpeta '{NOMBRE_CARPETA_RAIZ}' en mi Drive. "
        "Comprueba que existe y que el nombre coincide exactamente."
    )
ID_RAIZ = id_raiz["id"]
ID_CARPETA_A_PROCESAR = buscar_o_crear_carpeta("Documentos a procesar", ID_RAIZ)
ID_CARPETA_HISTORIAL = buscar_o_crear_carpeta("Historial Bonificaciones de Confirming", ID_RAIZ)


# Obtiene el banco, la sociedad y la fecha a partir del nombre del archivo.
def identificar_documento(ruta):
    """
    El formato esperado es BANCO_SOCIEDAD_FECHA, por ejemplo:
    'ORO_AZUL_01-01-26.pdf'.
    La separación por guiones bajos permite conocer
    estos datos antes de analizar el contenido del documento.
    """
    nombre_sin_extension, extension = os.path.splitext(os.path.basename(ruta))
    partes = nombre_sin_extension.split('_')
    banco = partes[0]
    sociedad = partes[1]
    fecha_documento = datetime.strptime(partes[2], '%d-%m-%y').date()
    return banco, sociedad, fecha_documento, extension.lower()

# PATRONES DE EXTRACCIÓN PARA LOS DOCUMENTOS BANCARIOS
"""
Cada banco presenta los datos con una distribución y unos formatos propios.
Por este motivo se definen expresiones regulares distintas para reconocer
CIF, facturas, importes, porcentajes y fechas. Estos patrones permiten
transformar documentos heterogéneos en una estructura común.
"""
import re

# Componentes comunes que se combinan para formar los patrones de cada banco.
CIF = r'[A-Z]\d{8}'
NUM_FACTURA = r'\d{10}'
IMPORTE_ES = r'-?[\d.]+,\d{2}'   # 1.234,56  (miles con punto, decimales con coma)
IMPORTE_US = r'-?[\d,]+\.\d{2}'  # 1,234.56  (miles con coma, decimales con punto)
PCT = r'\d{1,3},\d{2}'          # 12,34  (porcentaje con coma)
PCT_PUNTO = r'\d{1,3}\.\d{2}'    # 12.34  (porcentaje con punto, banco Rubi)
DIAS = r'\d{1,3}'

# Columnas procedentes del documento bancario que formarán la salida inicial.
COLUMNAS = ["CIF", "Proveedor", "Nº Factura", "Importe", "F. Anticipo",
            "F. Vencimiento", "Días", "% Interés", "% Comisión",
            "Total Comisión", "Total Interés", "% Reparto",
            "Comisión Prisma", "Interés Prisma"]

# Estos bancos presentan los catorce campos de cada factura en una sola línea.
REGEX_FILA = {
    "ORO": re.compile(fr'({CIF})\s+(.+?)\s+({NUM_FACTURA})\s+({IMPORTE_ES})\s+'
                       fr'(\d{{2}}-\d{{2}}-\d{{2}})\s+(\d{{2}}-\d{{2}}-\d{{2}})\s+({DIAS})\s+'
                       fr'({PCT})\s+({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                       fr'({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})'),

    "RUBI": re.compile(fr'({CIF})\s+(.+?)\s+({NUM_FACTURA})\s+({IMPORTE_US})\s+'
                        fr'(\d{{2}}/\d{{2}}/\d{{4}})\s+(\d{{2}}/\d{{2}}/\d{{4}})\s+({DIAS})\s+'
                        fr'({PCT_PUNTO})\s+({PCT_PUNTO})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                        fr'({PCT_PUNTO})\s+({IMPORTE_ES})\s+({IMPORTE_ES})'),

    "PERLA": re.compile(fr'({CIF})\s+(.+?)\s+({NUM_FACTURA})\s+({IMPORTE_US})\s+'
                         fr'(\d{{2}}\.\d{{2}}\.\d{{2}})\s+(\d{{2}}\.\d{{2}}\.\d{{2}})\s+({DIAS})\s+'
                         fr'({PCT})\s+({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                         fr'({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})'),

    "PLATINO": re.compile(fr'({CIF})\s+(.+?)\s+({NUM_FACTURA})\s+({IMPORTE_ES})\s+'
                           fr'(\d{{2}}/\d{{2}}/\d{{4}})\s+(\d{{2}}/\d{{2}}/\d{{4}})\s+({DIAS})\s+'
                           fr'({PCT})\s+({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                           fr'({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})'),
}

# Plata y Diamante muestran una cabecera con el CIF y el proveedor.
# Las facturas siguientes pertenecen a esa cabecera hasta que aparece otra.
REGEX_CABECERA = {
    "PLATA": re.compile(fr'CIF:\s*({CIF})\s+Proveedor:\s*(.+?)\s+Nº Factura'),
    "DIAMANTE": re.compile(fr'C\.I\.F\.:\s*({CIF})\s+Proveedor:\s*(.+?)\s+Nº Factura'),
}
REGEX_FILA_SIN_CIF = {
    "PLATA": re.compile(fr'({NUM_FACTURA})\s+({IMPORTE_ES})\s+'
                         fr'(\d{{2}}-\d{{2}}-\d{{4}})\s+(\d{{2}}-\d{{2}}-\d{{4}})\s+({DIAS})\s+'
                         fr'({PCT})\s+({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                         fr'({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})'),
    "DIAMANTE": re.compile(fr'({NUM_FACTURA})\s+({IMPORTE_ES})\s+'
                            fr'(\d{{2}}-\d{{2}}-\d{{2}})\s+(\d{{2}}-\d{{2}}-\d{{2}})\s+({DIAS})\s+'
                            fr'({PCT})\s+({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                            fr'({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})'),
}

# Esmeralda se recibe como una imagen escaneada y contiene dos tablas.
# Un patrón extrae la identificación y otro recupera los datos económicos.
REGEX_ESMERALDA_A = re.compile(fr'^(\d{{1,3}})\s+({CIF})\s+(.+?)\s+({NUM_FACTURA})\s+'
                                fr'(\d{{2}}\.\d{{2}}\.\d{{4}})\s+(\d{{2}}\.\d{{2}}\.\d{{4}})$', re.MULTILINE)
REGEX_ESMERALDA_B = re.compile(fr'^(\d{{1,3}})\s+({IMPORTE_ES})\s+({DIAS})\s+'
                                fr'({PCT})\s+({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})\s+'
                                fr'({PCT})\s+({IMPORTE_ES})\s+({IMPORTE_ES})$', re.MULTILINE)

print("Regex definidos para los 7 bancos en PDF")


# Relación entre cada banco y el formato de fecha presente en sus documentos.
FORMATO_FECHA = {
    "ORO": "%d-%m-%y", "DIAMANTE": "%d-%m-%y",
    "PLATA": "%d-%m-%Y",
    "RUBI": "%d/%m/%Y", "PLATINO": "%d/%m/%Y",
    "PERLA": "%d.%m.%y",
    "ESMERALDA": "%d.%m.%Y",
}


# Convierte un importe o porcentaje extraído como texto en un número.
def a_numero(texto):
    """
    La función reconoce si el último separador es la coma o el punto decimal
    y elimina el separador de miles cuando existe. Así admite valores escritos
    tanto con formato español como con formato anglosajón.
    """
    texto = texto.strip()
    if ',' in texto and '.' in texto:
        # El último separador indica cuál de los dos actúa como decimal.
        if texto.rfind(',') > texto.rfind('.'):
            texto = texto.replace('.', '').replace(',', '.')  # Convierte el formato español.
        else:
            texto = texto.replace(',', '')                     # Convierte el formato anglosajón.
    elif ',' in texto:
        texto = texto.replace(',', '.')  # Una coma aislada se interpreta como decimal.
    return float(texto)


# Convierte una fecha textual en un objeto de fecha de Python.
def a_fecha(texto, banco):
    """
    El formato adecuado se obtiene del diccionario FORMATO_FECHA mediante
    el nombre del banco. De esta forma, una sola función interpreta las
    distintas formas de escritura presentes en los documentos.
    """
    return datetime.strptime(texto, FORMATO_FECHA[banco]).date()


# Asigna a cada campo de una factura el tipo de dato adecuado.
def convertir_fila(fila_textual, banco):
    """
    La fila llega como una colección de textos obtenida mediante una expresión
    regular. La función conserva los campos descriptivos y transforma importes,
    porcentajes, días y fechas antes de incorporarlos al archivo de salida.
    """
    cif, proveedor, factura, importe, f_ant, f_ven, dias, pint, pcom, tcom, tint, prep, cprisma, iprisma = fila_textual
    return [
        cif, proveedor, factura,                    # Estos tres campos se conservan como texto.
        a_numero(importe),
        a_fecha(f_ant, banco), a_fecha(f_ven, banco),
        int(dias),
        a_numero(pint), a_numero(pcom),
        a_numero(tcom), a_numero(tint),
        a_numero(prep),
        a_numero(cprisma), a_numero(iprisma),
    ]


# Extrae las facturas del documento descargado en la máquina temporal.
def extraer_filas(ruta, banco, extension):
    """
    La función diferencia entre el Excel de Zafiro y los PDF del resto de
    bancos. En los PDF selecciona el patrón correspondiente y, cuando el
    documento está escaneado, aplica OCR antes de convertir los resultados.
    """
    if extension == '.xlsx':
        # Zafiro entrega un Excel cuyos datos comienzan en la cuarta fila.
        # Se seleccionan únicamente las columnas necesarias para el proceso.
        libro_original = openpyxl.load_workbook(ruta)
        hoja_original = libro_original.active
        columnas_utiles = [1, 2, 4, 5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]  # Se omiten las columnas C y F.
        return [[fila[col - 1].value for col in columnas_utiles]
                for fila in hoja_original.iter_rows(min_row=4)]

    # En los PDF se intenta primero extraer el texto que contiene cada página.
    from pypdf import PdfReader
    lector = PdfReader(ruta)
    texto = " ".join(pagina.extract_text() for pagina in lector.pages)

    # Si el PDF de Esmeralda no contiene texto suficiente, se considera escaneado.
    # Cada página se convierte en PNG y Tesseract aplica reconocimiento óptico
    # de caracteres en español para reconstruir el contenido del documento.
    if banco == "ESMERALDA" and len(texto.strip()) < 20:
        import subprocess
        import pytesseract
        from PIL import Image

        subprocess.run(["pdftoppm", "-png", "-r", "200", ruta, "pagina"], check=True)
        paginas_png = sorted(f for f in os.listdir('.') if f.startswith('pagina') and f.endswith('.png'))
        texto = ""
        for pagina in paginas_png:
            texto += pytesseract.image_to_string(Image.open(pagina), lang='spa') + "\n"
            os.remove(pagina)  # Elimina la imagen utilizada como archivo intermedio.

    filas_brutas = []
    if banco == "ESMERALDA":
        # El texto obtenido por OCR conserva una estructura aproximada por filas.
        filas_identificacion = REGEX_ESMERALDA_A.findall(texto)
        filas_detalle = REGEX_ESMERALDA_B.findall(texto)
        for ident, detalle in zip(filas_identificacion, filas_detalle):
            _, cif, proveedor, factura, f_ant, f_ven = ident
            _, importe, dias, pint, pcom, tcom, tint, prep, cprisma, iprisma = detalle
            fila = [cif, proveedor, factura, importe, f_ant, f_ven,
                    dias, pint, pcom, tcom, tint, prep, cprisma, iprisma]
            filas_brutas.append(convertir_fila(fila, banco))

    elif banco in REGEX_CABECERA:
        # En Plata y Diamante se combinan las cabeceras con sus facturas.
        # El CIF y el proveedor actuales se mantienen hasta encontrar una
        # nueva cabecera dentro del texto normalizado.
        texto_plano = re.sub(r'\s+', ' ', texto)
        patron_combinado = re.compile(
            f'(?:{REGEX_CABECERA[banco].pattern})|(?:{REGEX_FILA_SIN_CIF[banco].pattern})'
        )
        cif_actual, proveedor_actual = None, None
        for coincidencia in patron_combinado.finditer(texto_plano):
            grupos = coincidencia.groups()
            if grupos[0] is not None:
                cif_actual, proveedor_actual = grupos[0], grupos[1]
            else:
                fila = [cif_actual, proveedor_actual] + list(grupos[2:])
                filas_brutas.append(convertir_fila(fila, banco))

    else:
        # En los demás bancos, cada coincidencia contiene los catorce campos.
        texto_plano = re.sub(r'\s+', ' ', texto)
        for fila in REGEX_FILA[banco].findall(texto_plano):
            filas_brutas.append(convertir_fila(list(fila), banco))

    return filas_brutas

# CONSTRUCCIÓN DE LOS ARCHIVOS DE SALIDA
"""
Esta parte crea las hojas "Detalle" e "Imprimir" del libro de Excel.
También incorpora las columnas calculadas, aplica formatos y añade una
fila de totales. La hoja "Imprimir" servirá posteriormente como origen
del PDF entregado para cada documento procesado.
"""
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

# Columnas calculadas por el programa a partir de los datos bancarios.
COLUMNAS_EXTRA = ["Calculo com Prisma", "Calculo int prisma", "Total diferencia"]
COLUMNAS_BASE = COLUMNAS + COLUMNAS_EXTRA                                  # Conjunto de diecisiete columnas.
COLUMNAS_DETALLE = ["Banco", "Sociedad", "Fecha Documento"] + COLUMNAS_BASE  # Añade datos de identificación.
COLUMNAS_IMPRIMIR = [c for c in COLUMNAS_BASE if c not in ("Proveedor", "Nº Factura")]  # Excluye datos no mostrados en el PDF.

# Columnas económicas que se formatean y se suman en la fila de totales.
COLUMNAS_SUMA = ["Importe", "Total Comisión", "Total Interés", "Comisión Prisma",
                  "Interés Prisma", "Calculo com Prisma", "Calculo int prisma", "Total diferencia"]

# Columnas que se mostrarán con formato de fecha y año de dos cifras.
COLUMNAS_FECHA = ["Fecha Documento", "F. Anticipo", "F. Vencimiento"]

# Saltos de línea definidos para mejorar la presentación de títulos largos.
SALTOS_TITULO = {
    "Calculo com Prisma": "Calculo\ncom Prisma",
    "Calculo int prisma": "Calculo\nint prisma",
    "Total diferencia": "Total\ndiferencia",
}


# Construye el libro de Excel con las hojas "Detalle" e "Imprimir".
def construir_libro(banco, sociedad, fecha_documento, filas_brutas):
    """
    La primera hoja conserva todos los campos y añade banco, sociedad y fecha.
    La segunda excluye proveedor y número de factura porque está preparada para
    convertirse en PDF. Las fórmulas se incorporan en la fase de formato.
    """
    banco_legible = banco.title()
    sociedad_legible = sociedad.title()

    libro = openpyxl.Workbook()
    hoja_detalle = libro.active
    hoja_detalle.title = "Detalle"
    hoja_imprimir = libro.create_sheet("Imprimir")
    hoja_detalle.append(COLUMNAS_DETALLE)
    hoja_imprimir.append(COLUMNAS_IMPRIMIR)

    for fila in filas_brutas:
        cif, proveedor, factura, importe, f_ant, f_ven, dias, pint, pcom, tcom, tint, prep, cprisma, iprisma = fila
        hoja_detalle.append([banco_legible, sociedad_legible, fecha_documento,
                              cif, proveedor, factura, importe, f_ant, f_ven, dias, pint, pcom,
                              tcom, tint, prep, cprisma, iprisma])
        hoja_imprimir.append([cif, importe, f_ant, f_ven, dias, pint, pcom,
                               tcom, tint, prep, cprisma, iprisma])
    return libro, hoja_detalle, hoja_imprimir


# Aplica las fórmulas y el formato final a una hoja del libro.
def formatear_y_convertir_en_tabla(hoja, nombre_tabla):
    """
    La función ajusta tipos de datos, alineación, anchura de columnas y
    configuración de impresión. Finalmente convierte el rango en una tabla
    de Excel con una fila de totales, tanto en "Detalle" como en "Imprimir".
    """
    encabezados = [celda.value for celda in hoja[1]]
    columna = {nombre: i + 1 for i, nombre in enumerate(encabezados)}  # Posición de cada columna desde el valor 1.
    n_filas_datos = hoja.max_row  # Número de filas anterior a la incorporación de los totales.

    # Devuelve la letra de Excel asociada al nombre de una columna.
    def letra(nombre_columna):
        return hoja.cell(row=1, column=columna[nombre_columna]).column_letter

    # Las fórmulas se escriben con la letra real de cada columna.
    # Esto permite utilizar el mismo bloque aunque las columnas ocupen
    # posiciones diferentes en las hojas "Detalle" e "Imprimir".
    for r in range(2, n_filas_datos + 1):
        importe = f"{letra('Importe')}{r}"
        pcom = f"{letra('% Comisión')}{r}"
        preparto = f"{letra('% Reparto')}{r}"
        pint = f"{letra('% Interés')}{r}"
        dias = f"{letra('Días')}{r}"
        cprisma = f"{letra('Comisión Prisma')}{r}"
        iprisma = f"{letra('Interés Prisma')}{r}"
        ccom = f"{letra('Calculo com Prisma')}{r}"
        cint = f"{letra('Calculo int prisma')}{r}"

        hoja.cell(row=r, column=columna["Calculo com Prisma"],
                  value=f"={importe}*({pcom}/100)*({preparto}/100)")
        hoja.cell(row=r, column=columna["Calculo int prisma"],
                  value=f"={importe}*({pint}/360/100)*{dias}*({preparto}/100)")
        hoja.cell(row=r, column=columna["Total diferencia"],
                  value=f"={cprisma}-{ccom}+{iprisma}-{cint}")

    fuente_normal = Font(size=8)
    fuente_negrita = Font(size=8, bold=True)
    centrado = Alignment(horizontal='center', vertical='center')

    # Cada fila recibe el formato correspondiente a su tipo de información.
    # Los números de factura pasan a enteros, las fechas adoptan un formato
    # uniforme y los importes muestran separadores de miles y dos decimales.
    for fila in hoja.iter_rows(min_row=2, max_row=n_filas_datos):
        if "Nº Factura" in columna:
            fila[columna["Nº Factura"] - 1].value = int(fila[columna["Nº Factura"] - 1].value)
        for nombre_col in COLUMNAS_FECHA:
            if nombre_col in columna:
                fila[columna[nombre_col] - 1].number_format = 'dd/mm/yy'
        for nombre_col in COLUMNAS_SUMA:
            fila[columna[nombre_col] - 1].number_format = '#,##0.00'
        for celda in fila:
            celda.font = fuente_normal
            celda.alignment = centrado

    # La cabecera se presenta en negrita, centrada y con ajuste de texto.
    # Los títulos más largos emplean saltos definidos para facilitar su lectura.
    for celda in hoja[1]:
        if celda.value in SALTOS_TITULO:
            celda.value = SALTOS_TITULO[celda.value]
        celda.font = fuente_negrita
        celda.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    hoja.row_dimensions[1].height = 26

    # El ancho se calcula a partir del encabezado y del contenido visible.
    # Cuando una celda contiene una fórmula, se utiliza como referencia el
    # tamaño habitual de un importe para evitar columnas excesivamente anchas.
    for col_celdas in hoja.columns:
        letra_columna = col_celdas[0].column_letter
        texto_cabecera = col_celdas[0].value
        ancho_maximo = max(len(palabra) for linea in texto_cabecera.split(chr(10)) for palabra in linea.split())
        for celda in col_celdas[1:1 + n_filas_datos]:
            valor = celda.value
            if hasattr(valor, 'strftime'):
                texto = valor.strftime('%d/%m/%y')
            elif isinstance(valor, str) and valor.startswith('='):
                texto = "0.000,00"  # Representación utilizada para calcular el ancho de una fórmula.
            elif isinstance(valor, (int, float)):
                texto = f"{valor:,.2f}"
            else:
                texto = str(valor)
            ancho_maximo = max(ancho_maximo, len(texto))
        hoja.column_dimensions[letra_columna].width = ancho_maximo + 1

    # Las columnas calculadas mantienen un ancho fijo porque sus títulos ocupan dos líneas.
    for nombre_col in COLUMNAS_EXTRA:
        if nombre_col in columna:
            letra_columna = hoja.cell(row=1, column=columna[nombre_col]).column_letter
            hoja.column_dimensions[letra_columna].width = 11

    # La página se prepara en horizontal, con márgenes reducidos y cuadrícula.
    # La cabecera se repite al imprimir y el contenido se ajusta a una página de ancho.
    hoja.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3,
                                     header=0.1, footer=0.1)
    hoja.page_setup.orientation = 'landscape'
    hoja.print_options.gridLines = True
    hoja.print_title_rows = '1:1'
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True

    # El rango completo se convierte en una tabla de Excel sin estilo de color.
    # Los nombres de las columnas conservan el texto visible de la cabecera,
    # incluidos los saltos de línea, y se activa una fila final de totales.
    fila_totales = n_filas_datos + 1
    ultima_col_letra = hoja.cell(row=1, column=hoja.max_column).column_letter
    rango_tabla = f"A1:{ultima_col_letra}{fila_totales}"

    columnas_tabla = []
    for nombre_original, idx in sorted(columna.items(), key=lambda kv: kv[1]):
        texto_actual = hoja.cell(row=1, column=idx).value
        if idx == 1:
            columnas_tabla.append(TableColumn(id=idx, name=texto_actual, totalsRowLabel="Total"))
        elif nombre_original in COLUMNAS_SUMA:
            columnas_tabla.append(TableColumn(id=idx, name=texto_actual, totalsRowFunction="sum"))
        else:
            columnas_tabla.append(TableColumn(id=idx, name=texto_actual))

    tabla = Table(displayName=nombre_tabla, ref=rango_tabla, tableColumns=columnas_tabla, totalsRowCount=1)
    tabla.tableStyleInfo = TableStyleInfo(name=None, showRowStripes=False, showColumnStripes=False,
                                           showFirstColumn=False, showLastColumn=False)
    hoja.add_table(tabla)

    # Los totales se calculan mediante SUBTOTAL en lugar de valores fijos.
    # Por tanto, Excel puede recalcularlos si se modifica o filtra la tabla.
    hoja.cell(row=fila_totales, column=1, value="Total").font = fuente_negrita
    for nombre in COLUMNAS_SUMA:
        col_idx = columna[nombre]
        letra_col = hoja.cell(row=1, column=col_idx).column_letter
        formula_total = f"=SUBTOTAL(109,{letra_col}2:{letra_col}{n_filas_datos})"
        celda_total = hoja.cell(row=fila_totales, column=col_idx, value=formula_total)
        celda_total.number_format = '#,##0.00'
        celda_total.font = fuente_negrita
        celda_total.alignment = centrado


# Genera los archivos de salida y los guarda en Google Drive.
def procesar_y_guardar(banco, sociedad, fecha_documento, filas_brutas):
    """
    El libro se crea en /tmp, se formatean sus dos hojas y se obtiene el PDF
    mediante LibreOffice. Ambos archivos se suben a la carpeta de documentos
    procesados y la hoja "Detalle" se devuelve para actualizar el histórico.
    """
    libro, hoja_detalle, hoja_imprimir = construir_libro(banco, sociedad, fecha_documento, filas_brutas)
    formatear_y_convertir_en_tabla(hoja_detalle, "TablaDetalle")
    formatear_y_convertir_en_tabla(hoja_imprimir, "TablaImprimir")

    fecha_texto = fecha_documento.strftime('%d-%m-%y')
    nombre_base_salida = f"PROCESADO_{banco}_{sociedad}_{fecha_texto}"
    ruta_excel = f"/tmp/{nombre_base_salida}.xlsx"
    libro.save(ruta_excel)

    # Para crear el PDF se conserva únicamente la hoja "Imprimir".
    # LibreOffice realiza la conversión sin abrir una interfaz y utiliza la
    # configuración regional española para representar los decimales con coma.
    entorno = os.environ.copy()
    entorno["LANG"] = "es_ES.UTF-8"
    entorno["LC_ALL"] = "es_ES.UTF-8"

    libro_solo_imprimir = openpyxl.load_workbook(ruta_excel)
    for otra_hoja in list(libro_solo_imprimir.sheetnames):
        if otra_hoja != "Imprimir":
            del libro_solo_imprimir[otra_hoja]
    ruta_temporal = "/tmp/_tmp_imprimir.xlsx"
    libro_solo_imprimir.save(ruta_temporal)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", "/tmp", ruta_temporal],
                    check=True, env=entorno)
    ruta_pdf = f"/tmp/{nombre_base_salida}.pdf"
    os.replace("/tmp/_tmp_imprimir.pdf", ruta_pdf)
    os.remove(ruta_temporal)  # El archivo intermedio deja de ser necesario.

    # El Excel y el PDF se suben a la carpeta de la sociedad y el banco.
    # Después se eliminan las copias temporales de la máquina de GitHub.
    id_carpeta_procesados = buscar_o_crear_carpeta("Documentos Procesados", carpeta_banco(banco, sociedad))
    subir_archivo(ruta_excel, f"{nombre_base_salida}.xlsx", id_carpeta_procesados)
    subir_archivo(ruta_pdf, f"{nombre_base_salida}.pdf", id_carpeta_procesados)
    os.remove(ruta_excel)
    os.remove(ruta_pdf)

    print(f"  Excel y PDF subidos a Drive: {sociedad}/{banco}/Documentos Procesados")
    return hoja_detalle


# Prepara cada valor antes de incorporarlo al archivo histórico CSV.
def formatear_valor_csv(valor):
    """
    Las fechas se convierten a texto, las fórmulas se conservan y los números
    utilizan coma decimal. Como las columnas se separan mediante punto y coma,
    este formato evita confundir el separador decimal con el de los campos.
    """
    if hasattr(valor, "strftime"):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, str) and valor.startswith("="):
        return valor
    if isinstance(valor, float):
        return str(valor).replace(".", ",")
    return valor


# Añade las nuevas facturas al histórico acumulado en formato CSV.
def actualizar_historico(hoja_detalle):
    """
    Si el archivo ya existe en Drive, se descarga y se amplía con las filas
    de la hoja "Detalle". Después se actualiza el mismo archivo para mantener
    un único histórico; si no existe, se crea durante la primera ejecución.
    """
    nombre_csv = "Historial Bonificaciones de Confirming.csv"
    archivos_historial = listar_archivos(ID_CARPETA_HISTORIAL)
    archivo_existente = next((a for a in archivos_historial if a["title"] == nombre_csv), None)
    ya_existia_csv = archivo_existente is not None

    ruta_local_csv = f"/tmp/{nombre_csv}"
    if ya_existia_csv:
        archivo_existente.GetContentFile(ruta_local_csv)

    ultima_fila_con_datos = hoja_detalle.max_row - 1  # Excluye la fila de totales.
    primera_fila_a_copiar = 2 if ya_existia_csv else 1  # Evita repetir la cabecera en un histórico existente.

    # La codificación utf-8-sig facilita que Excel y Sheets reconozcan las tildes.
    # Puede utilizarse tanto al crear el CSV como al añadir nuevas filas.
    with open(ruta_local_csv, "a", newline="", encoding="utf-8-sig") as archivo_csv:
        escritor = csv.writer(archivo_csv, delimiter=";")
        for numero_fila, fila in enumerate(hoja_detalle.iter_rows(min_row=primera_fila_a_copiar, max_row=ultima_fila_con_datos, values_only=True), start=primera_fila_a_copiar):
            if numero_fila == 1:
                # La cabecera del CSV sustituye los saltos de línea por espacios.
                fila_texto = [valor.replace("\n", " ") if isinstance(valor, str) else valor for valor in fila]
            else:
                fila_texto = [formatear_valor_csv(valor) for valor in fila]
            escritor.writerow(fila_texto)

    # El histórico existente se actualiza con el mismo id y conserva su enlace.
    if ya_existia_csv:
        archivo_existente.SetContentFile(ruta_local_csv)
        archivo_existente.Upload()
    else:
        subir_archivo(ruta_local_csv, nombre_csv, ID_CARPETA_HISTORIAL)
    os.remove(ruta_local_csv)

    print("  Histórico actualizado:", nombre_csv)


# Archiva el documento original después de completar su tratamiento.
def archivar_documento(archivo_drive, banco, sociedad):
    """
    El archivo se mueve a la carpeta "Documentos Originales" del banco y la
    sociedad correspondientes. Si la estructura todavía no existe, se crea
    antes del traslado para impedir que el documento vuelva a procesarse.
    """
    id_carpeta_originales = buscar_o_crear_carpeta("Documentos Originales", carpeta_banco(banco, sociedad))
    mover_archivo(archivo_drive, id_carpeta_originales)
    print(f"  Archivado en Drive: {sociedad}/{banco}/Documentos Originales/{archivo_drive['title']}")


# Coordina todas las operaciones necesarias para procesar un documento.
def procesar_documento(archivo_drive):
    """
    El archivo se descarga, identifica y transforma antes de generar las
    salidas, actualizar el histórico y archivar el original. Si alguna fase
    falla, el error se registra y la identificación sigue formando parte del informe.
    """
    nombre_archivo = archivo_drive["title"]
    ruta_local = f"/tmp/{nombre_archivo}"
    archivo_drive.GetContentFile(ruta_local)
    banco, sociedad, fecha_documento, extension = identificar_documento(ruta_local)
    nombre_mostrar = f"{banco}_{sociedad}_{fecha_documento.strftime('%d-%m-%y')}"
    print(f"Procesando {nombre_archivo} (banco={banco}, sociedad={sociedad})")

    try:
        filas_brutas = extraer_filas(ruta_local, banco, extension)
        hoja_detalle = procesar_y_guardar(banco, sociedad, fecha_documento, filas_brutas)
        actualizar_historico(hoja_detalle)
        archivar_documento(archivo_drive, banco, sociedad)
    except Exception as error:
        print(f"AVISO: fallo al procesar {nombre_archivo}: {error}")
    finally:
        os.remove(ruta_local)  # Elimina siempre la copia local descargada.

    return banco, sociedad, nombre_mostrar


# Inicia el tratamiento de todos los documentos pendientes.
def procesar_todos_los_documentos_pendientes():
    """
    La función selecciona archivos PDF y Excel y los procesa uno a uno. Un
    error individual se muestra sin detener el resto del lote. Al finalizar
    genera el PDF de monitorización, aunque la carpeta de entrada esté vacía.
    """
    pendientes = listar_archivos(ID_CARPETA_A_PROCESAR)
    pendientes = [a for a in pendientes if a["title"].lower().endswith((".pdf", ".xlsx"))]

    subidos = []  # Guarda banco, sociedad y nombre de cada documento recibido.
    for archivo_drive in pendientes:
        try:
            subidos.append(procesar_documento(archivo_drive))
        except Exception as error:
            print(f"AVISO: no se ha podido ni identificar {archivo_drive['title']}: {error}")

    print("Proceso terminado." if pendientes else "No hay documentos pendientes en 'Documentos a procesar'.")

    generar_pdf_resumen(subidos)

# CREACIÓN DEL PDF DE MONITORIZACIÓN
# El proceso espera un documento por cada combinación de ocho sociedades y
# ocho bancos, es decir, un total habitual de 64 documentos. El informe indica
# si la ejecución cumple esa distribución y detalla los documentos recibidos
# para facilitar la revisión de ausencias o duplicidades.
SOCIEDADES = ("AMARILLO", "AZUL", "BLANCO", "ESCARLATA", "NEGRO", "PURPURA", "ROJO", "VERDE")
BANCOS = ("DIAMANTE", "ESMERALDA", "ORO", "PERLA", "PLATA", "PLATINO", "RUBI", "ZAFIRO")


# Organiza la información que aparecerá en el PDF de monitorización.
def preparar_resumen(subidos):
    """
    Los documentos se agrupan por sociedad y banco y se comparan con la cantidad
    esperada. A partir de este resultado se determina si la ejecución ha sido
    habitual o extraordinaria y se redacta la cabecera del informe.
    """
    por_combinacion = {}
    for banco, sociedad, nombre_mostrar in subidos:
        por_combinacion.setdefault((sociedad, banco), []).append(nombre_mostrar)

    total_esperado = len(SOCIEDADES) * len(BANCOS)
    cantidad_recibida = len(subidos)
    unidad = "documento" if cantidad_recibida == 1 else "documentos"
    es_habitual = all(len(por_combinacion.get((sociedad, banco), [])) == 1
                       for sociedad in SOCIEDADES for banco in BANCOS)

    if es_habitual:
        cabecera = (f"Documentos procesados de forma HABITUAL, se ha subido 1 documento para cada "
                    f"banco en cada sociedad. En total se han procesado {cantidad_recibida} {unidad} "
                    f"cuando lo habitual son {total_esperado} documentos.")
    else:
        cabecera = (f"Documentos procesados de forma EXTRAORDINARIA, revisar los archivos que se han "
                    f"subido para cada banco en cada sociedad. En total se han procesado "
                    f"{cantidad_recibida} {unidad} cuando lo habitual son {total_esperado} documentos.")

    return es_habitual, cabecera, por_combinacion


# Genera el informe de la ejecución y lo guarda en Google Drive.
def generar_pdf_resumen(subidos):
    """
    El PDF incluye la fecha y hora de ejecución, el resultado general y el
    detalle de cada sociedad y banco. Se guarda en "Historial Bonificaciones
    de Confirming" con el nombre "Documentos Procesados [fecha].pdf".
    """
    fecha_ejecucion = datetime.now(ZoneInfo("Europe/Madrid"))
    fecha_nombre = fecha_ejecucion.strftime("%d-%m-%Y")
    nombre_pdf = f"Documentos Procesados {fecha_nombre}.pdf"
    ruta_pdf = f"/tmp/{nombre_pdf}"
    es_habitual, cabecera, por_combinacion = preparar_resumen(subidos)
    print(f"  Generando PDF de monitorización con {len(subidos)} documento(s) recibido(s).")

    # Define los estilos que se utilizarán en las distintas partes del informe.
    estilos_base = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "TituloInforme",
        parent=estilos_base["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    estilo_estado = ParagraphStyle(
        "EstadoInforme",
        parent=estilos_base["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        spaceAfter=10,
    )
    estilo_sociedad = ParagraphStyle(
        "SociedadInforme",
        parent=estilos_base["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=12,
        spaceBefore=5,
        spaceAfter=2,
    )
    estilo_linea = ParagraphStyle(
        "LineaInforme",
        parent=estilos_base["BodyText"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=10,
        leftIndent=5 * mm,
        spaceAfter=1,
    )

    # Prepara la cabecera del PDF con la fecha, la hora y el resultado general.
    elementos = [
        Paragraph(f"Documentos Procesados {fecha_nombre}", estilo_titulo),
        Paragraph(
            "Ejecución de GitHub Actions: "
            + fecha_ejecucion.strftime("%d/%m/%Y %H:%M, hora de España"),
            estilos_base["BodyText"],
        ),
        Spacer(1, 5 * mm),
        Paragraph(escape(cabecera), estilo_estado),
    ]

    for sociedad in SOCIEDADES:
        # Cada bloque mantiene unidos el nombre de la sociedad y sus ocho bancos.
        bloque_sociedad = [Paragraph(f"Sociedad {sociedad.title()}", estilo_sociedad)]
        for banco in BANCOS:
            nombres = por_combinacion.get((sociedad, banco), [])
            valor = ", ".join(nombres) if nombres else "(no se ha recibido ningún documento)"
            bloque_sociedad.append(Paragraph(escape(f"Banco {banco.title()}: {valor}"), estilo_linea))
        elementos.append(KeepTogether(bloque_sociedad))

    def pie_pagina(lienzo, documento):
        # Añade el número de página en el margen inferior del informe.
        lienzo.saveState()
        lienzo.setFont("Helvetica", 8)
        lienzo.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Página {documento.page}")
        lienzo.restoreState()

    # Configura el tamaño de página, los márgenes y los metadatos del PDF.
    documento = SimpleDocTemplate(
        ruta_pdf,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Bonificaciones confirming - proceso {'habitual' if es_habitual else 'extraordinario'}",
        author="Pipeline de auditoría de confirming",
    )
    documento.build(elementos, onFirstPage=pie_pagina, onLaterPages=pie_pagina)

    subir_o_actualizar_archivo(ruta_pdf, nombre_pdf, ID_CARPETA_HISTORIAL)
    os.remove(ruta_pdf)
    print(f"  PDF de monitorización guardado: Historial Bonificaciones de Confirming/{nombre_pdf}")

# INICIO DEL PROGRAMA Y PLANIFICACIÓN EXTERNA
"""
La fecha de ejecución no se controla desde Python, sino mediante el cron
definido en mensual.yml. Cada vez que GitHub inicia el programa, este procesa
todos los documentos pendientes, tanto en una ejecución programada como al
utilizar manualmente el botón "Run workflow".
"""
if __name__ == "__main__": # así ejecuta Github Actions el programa
    procesar_todos_los_documentos_pendientes()
